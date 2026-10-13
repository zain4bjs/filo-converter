"""Unit tests for watermark/filo_watermark.py.

    python -m pytest watermark/tests   (PyMuPDF, python-docx, python-pptx,
                                        openpyxl, Pillow: the converter pins)
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import filo_watermark as fw  # noqa: E402


def test_header_value():
    assert fw.watermark_requested("1") and fw.watermark_requested(" 1 ")
    for value in (None, "", "0", "true", "yes", "11", "1,0"):
        assert not fw.watermark_requested(value), value


def test_assets_are_transparent():
    from PIL import Image

    for path in (fw.LOGO_PLAIN, fw.LOGO_MARK):
        with Image.open(path) as image:
            assert image.mode == "RGBA"
            assert image.getpixel((0, 0))[3] == 0
    with Image.open(fw.LOGO_MARK) as image:
        assert max(image.getchannel("A").getdata()) <= round(255 * fw.OPACITY) + 1


def test_mark_width_clamped():
    assert fw.mark_width(595, 842) == pytest.approx(595 * 0.14)
    assert fw.mark_width(5000, 5000) == fw.MAX_WIDTH_PT
    assert fw.mark_width(300, 400) == fw.MIN_WIDTH_PT
    assert fw.mark_width(100, 100) <= 40


def test_pdf_every_page(tmp_path):
    import fitz

    doc = fitz.open()
    for rotation in (0, 90, 180, 270):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 72), "hello")
        page.set_rotation(rotation)
    path = tmp_path / "in.pdf"
    doc.save(path)
    assert fw.watermark_pdf(path) == 4
    with fitz.open(path) as out:
        for page in out:
            assert "filodocs.app" in page.get_text() and "Made with Filo" in page.get_text()
            assert len(page.get_images()) == 1
            # The text sits in the bottom-right quarter of the visible page.
            hit = page.search_for("filodocs.app")[0] * page.rotation_matrix
            assert hit.x0 > page.rect.width / 2 and hit.y0 > page.rect.height * 0.8, page.rotation
        # One image object shared by every page.
        assert len({img[0] for page in out for img in page.get_images()}) == 1



@pytest.mark.parametrize("rotation", (0, 90, 180, 270))
@pytest.mark.parametrize(
    "mediabox,cropbox",
    (
        ((0, 0, 800, 1000), (100, 200, 500, 700)),
        ((0, 0, 800, 1000), (0, 0, 400, 500)),
        ((-50, -80, 562, 762), None),
        ((30, 40, 640, 900), (100, 200, 500, 700)),
    ),
)
def test_pdf_logo_sits_above_text_on_offset_and_rotated_pages(tmp_path, mediabox, cropbox, rotation):
    """The wordmark picture stays on the page, bottom right, centred over the
    two text lines, whatever the media/crop box origin and /Rotate."""
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=mediabox[2] - mediabox[0], height=mediabox[3] - mediabox[1])
    if mediabox[0] or mediabox[1]:
        page.set_mediabox(fitz.Rect(*mediabox))
    if cropbox:
        page.set_cropbox(fitz.Rect(*cropbox))
    page.set_rotation(rotation)
    path = tmp_path / "in.pdf"
    doc.save(path)
    fw.watermark_pdf(path)
    with fitz.open(path) as out:
        pix = out[0].get_pixmap(dpi=72)
    width, height, samples, n = pix.width, pix.height, pix.samples, pix.n
    logo, text = [], []
    for y in range(height):
        for x in range(width):
            r, g, b = samples[(y * width + x) * n:(y * width + x) * n + 3]
            if r > 245 and g > 245 and b > 245:
                continue
            (logo if b > r + 40 else text).append((x, y))
    assert logo and text, "logo or text missing from the visible page"
    lx0, ly0 = min(p[0] for p in logo), min(p[1] for p in logo)
    lx1, ly1 = max(p[0] for p in logo), max(p[1] for p in logo)
    below = [p for p in text if p[1] >= ly1 - 1]
    tx0, ty0, tx1 = min(p[0] for p in below), min(p[1] for p in below), max(p[0] for p in below)
    assert lx0 > width / 2 and ly0 > height / 2 and lx1 < width - 3
    assert abs((lx0 + lx1) / 2 - (tx0 + tx1) / 2) < 8
    assert -3 <= ty0 - ly1 <= 15

def test_docx_footer(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Body")
    document.add_section()
    document.sections[1].different_first_page_header_footer = True
    path = tmp_path / "in.docx"
    document.save(path)
    assert fw.watermark_docx(path) == 2  # section 1 default + section 2 first page
    out = Document(path)
    assert fw.FOOTER_TEXT in "".join(p.text for p in out.sections[0].footer.paragraphs)
    assert fw.FOOTER_TEXT in "".join(p.text for p in out.sections[1].first_page_footer.paragraphs)
    with zipfile.ZipFile(path) as archive:
        footers = [n for n in archive.namelist() if n.startswith("word/footer")]
        assert footers and all(b"<pic:pic" in archive.read(n) for n in footers)


def test_xlsx_footer(tmp_path):
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    workbook.active.append(["a", 1])
    workbook.create_sheet("Visuals")
    assert fw.watermark_workbook(workbook) == 2
    path = tmp_path / "out.xlsx"
    workbook.save(path)
    for sheet in load_workbook(path).worksheets:
        assert sheet.oddFooter.right.text == fw.FOOTER_TEXT


def test_pptx_every_slide(tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    for _ in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    assert fw.watermark_presentation(presentation) == 3
    path = tmp_path / "out.pptx"
    presentation.save(path)
    for slide in Presentation(path).slides:
        names = {shape.name for shape in slide.shapes}
        assert {"Filo watermark logo", "Filo watermark text"} <= names
        text = next(s for s in slide.shapes if s.name == "Filo watermark text").text_frame.text
        assert fw.LINE_1 in text and fw.LINE_2 in text
        logo = next(s for s in slide.shapes if s.name == "Filo watermark logo")
        assert logo.left > presentation.slide_width / 2 and logo.top > presentation.slide_height / 2
