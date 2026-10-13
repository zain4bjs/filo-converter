"""The "Made with Filo" watermark for converter outputs.

A converter adds it only when the request carries the header
X-Filo-Watermark: 1 (see watermark_requested and the README section
"Watermark"). In a deployment, a fronting proxy decides that value and
overwrites whatever a client sends. Absent or anything other than "1" means
no watermark.

The mark is a small block in
the bottom-right corner, inset from the edges, three centred lines:

    FILO wordmark      (assets/watermark-filo-mark.png: transparent, with a
                        soft white halo, already at the mark's 60 % opacity)
    Made with Filo
    filodocs.app

about 14 % of the page width (clamped), one per page or slide.

Per output type:
    PDF   every page: wordmark image + the two text lines (real text, white
          halo stroke under a navy fill, both at 60 % opacity)
    DOCX  every section's footers: wordmark image + "Made with Filo · filodocs.app",
          in a floating frame at the page's bottom right (no layout change)
    XLSX  every sheet's footer: the text "Made with Filo · filodocs.app"
          (images in sheet footers need VML; text is what Excel shows reliably)
    PPTX  every slide: wordmark picture + a two-line text box, bottom right
Text outputs (pdf-to-text) are never watermarked.

Libraries: PyMuPDF, python-docx, python-pptx, openpyxl and Pillow, all
already in the converter interpreter. Each is imported only by the function
that needs it.
"""

from __future__ import annotations

import os
from pathlib import Path

ASSETS = Path(__file__).resolve().parent / "assets"
LOGO_MARK = ASSETS / "watermark-filo-mark.png"
LOGO_PLAIN = ASSETS / "watermark-filo.png"

HEADER = "X-Filo-Watermark"
LINE_1 = "Made with Filo"
LINE_2 = "filodocs.app"
FOOTER_TEXT = f"{LINE_1} · {LINE_2}"

# Mark width as a share of the page width, clamped (PDF points).
WIDTH_SHARE = 0.14
MIN_WIDTH_PT = 56.0
MAX_WIDTH_PT = 150.0
OPACITY = 0.6
# Navy of the wordmark's "ILO" (#18254F) for text; white for the halo.
INK_RGB = (0x18, 0x25, 0x4F)
HALO_RGB = (0xFF, 0xFF, 0xFF)
# Footer text colour for DOCX/XLSX, where text opacity is not available.
FOOTER_GREY = "6B7280"


def watermark_requested(value: str | None) -> bool:
    """True only for the exact header value "1"."""
    return (value or "").strip() == "1"


def _logo_geometry() -> tuple[float, float]:
    """(halo image width / wordmark width, halo image height / width)."""
    from PIL import Image

    with Image.open(LOGO_MARK) as mark, Image.open(LOGO_PLAIN) as plain:
        return mark.width / plain.width, mark.height / mark.width


def mark_width(page_width: float, page_height: float) -> float:
    width = min(MAX_WIDTH_PT, max(MIN_WIDTH_PT, page_width * WIDTH_SHARE))
    # Tiny pages: never more than 40 % of either side.
    return min(width, 0.4 * page_width, 0.4 * page_height / 0.55)


# ----- PDF -------------------------------------------------------------------


def watermark_pdf(path: str | os.PathLike) -> int:
    """Stamp every page of the PDF at ``path`` in place; returns the page count."""
    import fitz  # PyMuPDF

    path = Path(path)
    overscan, aspect = _logo_geometry()
    logo_bytes = LOGO_MARK.read_bytes()
    ink = tuple(c / 255 for c in INK_RGB)
    halo = tuple(c / 255 for c in HALO_RGB)

    with fitz.open(path) as doc:
        if doc.needs_pass:
            raise ValueError("encrypted PDF")
        xref = 0
        for page in doc:
            visible = page.rect  # rotation-aware, origin at the visible top left
            width = mark_width(visible.width, visible.height)
            inset = max(8.0, min(36.0, 0.03 * min(visible.width, visible.height)))
            size_1 = width * 0.125
            size_2 = width * 0.11
            centre_x = visible.width - inset - width / 2
            baseline_2 = visible.height - inset - size_2 * 0.25
            baseline_1 = baseline_2 - size_2 * 1.3
            logo_bottom = baseline_1 - size_1 * 1.05
            image_w = width * overscan
            image_h = image_w * aspect
            # The halo image has (image_w - width) / 2 of halo padding round
            # the wordmark: place it so the wordmark itself ends at logo_bottom.
            image_rect = fitz.Rect(
                centre_x - image_w / 2,
                logo_bottom - image_h + (image_w - width) / 2,
                centre_x + image_w / 2,
                logo_bottom + (image_w - width) / 2,
            )
            derotate = page.derotation_matrix
            rotate = page.rotation
            # insert_image misplaces the picture on a rotated page whose crop
            # box or media box does not start at (0, 0). Place it on the
            # unrotated page (rect in unrotated coordinates, picture turned by
            # the page rotation), then restore the rotation.
            unrotated = (image_rect * derotate).normalize()
            if rotate:
                page.set_rotation(0)
            try:
                if xref:
                    page.insert_image(unrotated, xref=xref, rotate=rotate, keep_proportion=True, overlay=True)
                else:
                    xref = page.insert_image(
                        unrotated, stream=logo_bytes, rotate=rotate, keep_proportion=True, overlay=True
                    )
            finally:
                if rotate:
                    page.set_rotation(rotate)
            for text, size, font, baseline in (
                (LINE_1, size_1, "hebo", baseline_1),
                (LINE_2, size_2, "helv", baseline_2),
            ):
                text_w = fitz.get_text_length(text, fontname=font, fontsize=size)
                origin = fitz.Point(centre_x - text_w / 2, baseline) * derotate
                common = dict(fontname=font, fontsize=size, rotate=rotate, overlay=True)
                # Halo: a white stroke under the fill.
                page.insert_text(
                    origin, text, color=halo, render_mode=1, border_width=0.15,
                    stroke_opacity=OPACITY, fill_opacity=OPACITY, **common,
                )
                page.insert_text(
                    origin, text, color=ink, fill=ink, render_mode=0,
                    stroke_opacity=OPACITY, fill_opacity=OPACITY, **common,
                )
        count = doc.page_count
        temp = path.with_name(path.stem + ".wm" + path.suffix)
        doc.save(temp, garbage=1, deflate=True)
    temp.replace(path)
    return count


# ----- DOCX ------------------------------------------------------------------


def _docx_footers(document):
    """(section, footer, created) for every footer Word shows: the default
    one, plus first-page and even-page footers where the document uses them.
    A later section's footer that is linked to the previous section is
    skipped when that kind was stamped already (it shows the stamped one);
    otherwise it gets its own definition (created=True)."""
    stamped: set[str] = set()
    even = document.settings.odd_and_even_pages_header_footer
    for section in document.sections:
        kinds = [("default", section.footer)]
        if section.different_first_page_header_footer:
            kinds.append(("first", section.first_page_footer))
        if even:
            kinds.append(("even", section.even_page_footer))
        for kind, footer in kinds:
            linked = footer.is_linked_to_previous
            if linked and kind in stamped:
                continue
            stamped.add(kind)
            yield section, footer, linked


def _text_width_pt(text: str, size: float) -> float:
    try:
        import fitz

        return fitz.get_text_length(text, fontname="helv", fontsize=size)
    except Exception:
        return len(text) * size * 0.55


def _shrink_paragraph(paragraph) -> None:
    """Make an empty footer paragraph 1 pt tall, so a footer that only holds
    the (floating) mark takes almost no room from the page body."""
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(1)
    for run in paragraph.runs:
        run.font.size = Pt(1)
    paragraph_mark = paragraph._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_pr = paragraph_mark.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        paragraph_mark.append(r_pr)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "2")
        r_pr.append(el)


def watermark_docx(path: str | os.PathLike) -> int:
    """Add the mark to every footer of the .docx at ``path`` in place; returns
    the number of footers stamped.

    The mark is a framed paragraph (w:framePr) positioned on the page, bottom
    right: it floats, so it never takes room from the page body and never
    moves the converted layout (pdf2docx and LibreOffice place content to the
    point). Inside it, centred: the wordmark picture, then the footer text."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    overscan, aspect = _logo_geometry()
    path = Path(path)
    document = Document(str(path))
    stamped = 0
    for section, footer, created in _docx_footers(document):
        page_w = section.page_width.pt if section.page_width else 612.0
        page_h = section.page_height.pt if section.page_height else 792.0
        width = mark_width(page_w, page_h)
        inset = max(8.0, min(36.0, 0.03 * min(page_w, page_h)))
        text_size = max(6.0, min(9.0, width * 0.088))
        # Room to spare: Word and LibreOffice pad frames slightly and would
        # otherwise clip the picture's right edge.
        frame_w = max(width * overscan, _text_width_pt(FOOTER_TEXT, text_size)) * 1.15 + 6
        image_w = width * overscan
        image_h = image_w * aspect
        frame_h = image_h + text_size * 1.3
        if created:
            # The new footer's own empty paragraph: keep it 1 pt tall, and
            # keep the footer inside the bottom margin, so the body keeps
            # exactly the room it had without a footer.
            for paragraph in footer.paragraphs:
                if not paragraph.text.strip():
                    _shrink_paragraph(paragraph)
            bottom = section.bottom_margin.pt if section.bottom_margin is not None else 0.0
            distance = section.footer_distance.pt if section.footer_distance is not None else 0.0
            if distance > max(0.0, bottom - 2):
                section.footer_distance = Pt(max(0.0, bottom - 2))

        paragraph = footer.paragraphs[0].insert_paragraph_before() if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        logo = paragraph.add_run()
        logo.add_picture(str(LOGO_MARK), width=Pt(image_w))
        logo.add_break()
        text = paragraph.add_run(FOOTER_TEXT)
        text.font.size = Pt(text_size)
        text.font.name = "Arial"
        text.font.color.rgb = RGBColor.from_string(FOOTER_GREY)

        frame = OxmlElement("w:framePr")
        for name, value in (
            ("w:w", str(round(frame_w * 20))),
            ("w:hAnchor", "page"),
            ("w:vAnchor", "page"),
            ("w:x", str(round((page_w - inset - frame_w) * 20))),
            ("w:y", str(round((page_h - inset - frame_h) * 20))),
            ("w:wrap", "none"),
        ):
            frame.set(qn(name), value)
        p_pr = paragraph._p.get_or_add_pPr()
        style = p_pr.find(qn("w:pStyle"))
        if style is not None:
            style.addnext(frame)
        else:
            p_pr.insert(0, frame)
        stamped += 1
    document.save(str(path))
    return stamped


# ----- XLSX ------------------------------------------------------------------


def watermark_workbook(workbook) -> int:
    """Put the footer text on every sheet of an openpyxl Workbook (before it
    is saved); returns the number of sheets."""
    count = 0
    for sheet in workbook.worksheets:
        for footer in (sheet.oddFooter, sheet.evenFooter, sheet.firstFooter):
            footer.right.text = FOOTER_TEXT
            footer.right.size = 8
            footer.right.color = FOOTER_GREY
        count += 1
    return count


# ----- PPTX ------------------------------------------------------------------

_TEXT_GLOW = (
    '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:glow rad="25400"><a:srgbClr val="FFFFFF"><a:alpha val="{alpha}"/></a:srgbClr></a:glow>'
    "</a:effectLst>"
)


def watermark_presentation(presentation) -> int:
    """Add the mark to every slide of a python-pptx Presentation (before it
    is saved); returns the number of slides."""
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    overscan, aspect = _logo_geometry()
    emu_per_pt = 12700
    slide_w = presentation.slide_width / emu_per_pt
    slide_h = presentation.slide_height / emu_per_pt
    width = mark_width(slide_w, slide_h)
    inset = max(8.0, min(36.0, 0.03 * min(slide_w, slide_h)))
    size_1 = width * 0.125
    size_2 = width * 0.11
    text_h = size_1 * 1.25 + size_2 * 1.25
    centre_x = slide_w - inset - width / 2
    text_top = slide_h - inset - text_h
    image_w = width * overscan
    image_h = image_w * aspect
    pad = (image_w - width) / 2
    image_top = text_top - image_h + pad
    alpha = str(int(OPACITY * 100000))

    count = 0
    for slide in presentation.slides:
        picture = slide.shapes.add_picture(
            str(LOGO_MARK),
            Emu(int((centre_x - image_w / 2) * emu_per_pt)),
            Emu(int(image_top * emu_per_pt)),
            width=Emu(int(image_w * emu_per_pt)),
            height=Emu(int(image_h * emu_per_pt)),
        )
        picture.name = "Filo watermark logo"
        box = slide.shapes.add_textbox(
            Emu(int((centre_x - width * 0.75) * emu_per_pt)),
            Emu(int(text_top * emu_per_pt)),
            Emu(int(width * 1.5 * emu_per_pt)),
            Emu(int(text_h * emu_per_pt)),
        )
        box.name = "Filo watermark text"
        frame = box.text_frame
        frame.word_wrap = False
        frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(frame, margin, 0)
        for index, (text, size, bold) in enumerate(((LINE_1, size_1, True), (LINE_2, size_2, False))):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = text
            font = run.font
            font.size = Pt(round(size, 1))
            font.bold = bold
            font.color.rgb = RGBColor(*INK_RGB)
            r_pr = run._r.get_or_add_rPr()
            fill = r_pr.find(qn("a:solidFill"))
            colour = fill.find(qn("a:srgbClr"))
            etree.SubElement(colour, qn("a:alpha")).set("val", alpha)
            # effectLst must come after the fill in a:rPr.
            fill.addnext(etree.fromstring(_TEXT_GLOW.replace("{alpha}", "80000")))
        count += 1
    return count
