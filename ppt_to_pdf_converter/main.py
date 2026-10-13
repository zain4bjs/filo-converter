import asyncio
import atexit
import io
import logging
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import xml.etree.ElementTree as ElementTree
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from functools import lru_cache
from pathlib import Path
from uuid import uuid4

import fitz  # PyMuPDF
import pdfplumber
import PIL.ImageOps
import pytesseract
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, JSONResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
from starlette.background import BackgroundTask
from starlette.concurrency import run_in_threadpool
from starlette.exceptions import HTTPException as StarletteHTTPException

# Shared "Made with Filo" watermark (repo root watermark/, /app/watermark in
# the image). Applied only when the request carries X-Filo-Watermark: 1, and never
# to pdf-to-text output.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "watermark"))
import filo_watermark  # noqa: E402

# Per-request time budget (repo root deadline/, /app/deadline in the image):
# answers 504 before a typical 100 s upstream proxy limit. See filo_deadline.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "deadline"))
import filo_deadline  # noqa: E402
from filo_deadline import Deadline, DeadlineExceeded  # noqa: E402

# uvicorn only configures its own loggers; give ours a handler so server-side
# failure details (which are never sent to clients) actually reach the logs.
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("ppt_to_pdf_converter")


def _env_int(name: str, default: int) -> int:
    """Read a positive integer from the environment.

    Empty values and unexpanded CI macros ("$(VAR)") count as unset,
    and so does anything that does not parse: a typo must not take the service
    down or silently disable a limit.
    """
    raw = os.environ.get(name, "").strip()
    if not raw or raw.startswith("$("):
        return default
    try:
        value = int(raw)
    except ValueError:
        logger.warning("Ignoring invalid %s=%r; using %d", name, raw, default)
        return default
    return value if value > 0 else default


@asynccontextmanager
async def lifespan(app: FastAPI):
    """Build the LibreOffice profile template before the first request.

    Doing it here keeps the ~3s cost off whichever request happens to arrive
    first. A host without LibreOffice must still boot — the endpoints report
    the missing binary themselves — so a failure here is not fatal.
    """
    try:
        await run_in_threadpool(_profile_template)
    except Exception:
        logger.warning("LibreOffice profile template not built at startup", exc_info=True)
    yield


app = FastAPI(title="Office <-> PDF Converter", lifespan=lifespan)

# ----- Limits ----------------------------------------------------------------

# Uploads above this are refused with 413 (Content-Length pre-check, then a
# per-file check while streaming to disk).
MAX_UPLOAD_MB = _env_int("MAX_UPLOAD_MB", 30)
MAX_FILE_SIZE = MAX_UPLOAD_MB * 1024 * 1024
# A multipart body is the file plus boundaries and the other form fields, so
# the Content-Length pre-check leaves a little room for that envelope.
_MULTIPART_SLACK = 1024 * 1024

# Longer PDFs are refused with 422 before any rendering or extraction starts.
# A PDF under the cap that still cannot convert within the request budget
# (REQUEST_DEADLINE_SECONDS, below) gets a 504 instead (README, "Limits").
MAX_PDF_PAGES = _env_int("MAX_PDF_PAGES", 50)

# Total time a conversion request may take, counted from when a reverse proxy
# in front received it, or else from when this app did (filo_deadline):
# upload, queueing for a slot below and the conversion itself. 95 s by
# default, under a typical 100 s upstream proxy limit. Past it the route answers 504
# {"error": "Conversion took too long"}: a LibreOffice run is killed (its
# timeout is min(LO_TIMEOUT_SECONDS, what is left)), tesseract runs get the
# time left as their timeout, and the PyMuPDF / pdfplumber page loops stop at
# the next page.
REQUEST_DEADLINE_SECONDS = filo_deadline.budget_seconds()

# Concurrent LibreOffice runs and OCR jobs. Each soffice process takes a few
# hundred MB and a core; OCR fans out to OCR_WORKERS tesseract processes. Past
# this figure requests queue instead of exhausting the container's memory.
LO_MAX_CONCURRENCY = _env_int("LO_MAX_CONCURRENCY", 2)
_HEAVY_SLOTS = asyncio.Semaphore(LO_MAX_CONCURRENCY)

# Concurrent PyMuPDF / pdfplumber jobs (pdf-to-ppt page renders, pdf-to-xls
# table extraction, pdf-to-text text layers). Otherwise they only share
# Starlette's ~40-thread pool, and a handful of long PDFs at once could take
# every core and gigabytes of page renders.
PDF_MAX_CONCURRENCY = _env_int("PDF_MAX_CONCURRENCY", 2)
_PDF_SLOTS = asyncio.Semaphore(PDF_MAX_CONCURRENCY)

# Longest a request waits for one of the slots above before it gets a 503
# (which the app retries). The request budget above is the tighter bound in
# practice: a request still queued when its budget runs out gets the 504.
SLOT_WAIT_SECONDS = _env_int("SLOT_WAIT_SECONDS", 120)

# OOXML/ODF uploads are zip archives: 30 MB on the wire can unpack to many GB
# inside LibreOffice. Refused (422) when the declared unpacked size exceeds
# this. Real documents are far below it.
MAX_UNPACKED_MB = _env_int("MAX_UNPACKED_MB", 1000)

# Pillow refuses (DecompressionBombError) images over twice this many pixels.
Image.MAX_IMAGE_PIXELS = _env_int("MAX_IMAGE_PIXELS", 90_000_000)
# PyMuPDF has no such guard: a PDF can declare a page hundreds of inches wide,
# and rasterising it would allocate gigabytes. 150 MP still fits an A0 page at
# the 300 DPI OCR resolution (~139 MP).
MAX_RENDER_PIXELS = _env_int("MAX_RENDER_PIXELS", 150_000_000)


# ----- Error responses -------------------------------------------------------

# Clients only ever see a short fixed message under "error". Exception text,
# LibreOffice/Tesseract stderr, and file paths are logged here instead: they
# can leak server internals, and the app never displays them anyway.


def _deadline(request: Request) -> Deadline:
    """The request's budget, started by _GuardMiddleware when it arrived."""
    deadline = getattr(request.state, "deadline", None)
    if deadline is None:
        deadline = Deadline(REQUEST_DEADLINE_SECONDS)
        request.state.deadline = deadline
    return deadline


async def _acquire_slot(semaphore: asyncio.Semaphore, request: Request) -> None:
    """Take one slot of ``semaphore`` within SLOT_WAIT_SECONDS and the budget.

    Raises 504 when the request budget ran out while queued, 503 (retryable
    for the app) when SLOT_WAIT_SECONDS did. After getting the slot, work for
    a client that has already disconnected is skipped: Starlette does not
    cancel a request when its client goes away.
    """
    deadline = _deadline(request)
    wait = min(SLOT_WAIT_SECONDS, deadline.timeout())
    try:
        await asyncio.wait_for(semaphore.acquire(), wait)
    except asyncio.TimeoutError as exc:
        if deadline.expired():
            logger.warning("Request deadline ran out while queued on %s", request.url.path)
            raise DeadlineExceeded() from exc
        logger.warning("No free conversion slot after %ss on %s", SLOT_WAIT_SECONDS, request.url.path)
        raise HTTPException(status_code=503, detail="Server busy, please retry") from exc
    try:
        if await request.is_disconnected():
            logger.info("Client disconnected while queued on %s; skipping", request.url.path)
            raise HTTPException(status_code=503, detail="Client disconnected")
        deadline.check()
    except BaseException:
        semaphore.release()
        raise


async def _run_bounded(semaphore: asyncio.Semaphore, request: Request, func, *args):
    """Run blocking ``func(*args)`` in a worker thread holding one slot.

    Answers within the request budget: when it runs out first, the budget is
    marked expired (so ``func``'s own checks stop it at the next page, and a
    LibreOffice or tesseract run is killed by its own timeout at the same
    moment) and DeadlineExceeded (504) is raised at once. The slot is only
    released when the thread has really finished, so the concurrency caps
    still hold for work that is winding down.
    """
    deadline = _deadline(request)
    await _acquire_slot(semaphore, request)
    try:
        task = asyncio.ensure_future(run_in_threadpool(func, *args))
    except BaseException:
        semaphore.release()
        raise

    def _finished(done: asyncio.Future) -> None:
        semaphore.release()
        if not done.cancelled():
            done.exception()  # retrieved: no "exception never retrieved" noise

    task.add_done_callback(_finished)
    try:
        return await asyncio.wait_for(asyncio.shield(task), deadline.remaining())
    except asyncio.TimeoutError:
        if task.done():  # finished in the same instant
            return task.result()
        deadline.expire()
        logger.warning(
            "Request deadline (%.0f s) reached on %s; stopping the conversion",
            deadline.seconds,
            request.url.path,
        )
        raise DeadlineExceeded() from None


def _error(status_code: int, message: str, headers: dict | None = None) -> JSONResponse:
    return JSONResponse({"error": message}, status_code=status_code, headers=headers)


@app.exception_handler(StarletteHTTPException)
async def _http_error(request, exc: StarletteHTTPException) -> JSONResponse:
    message = exc.detail if isinstance(exc.detail, str) else "Request failed"
    return _error(exc.status_code, message, getattr(exc, "headers", None))


@app.exception_handler(DeadlineExceeded)
async def _deadline_error(request, exc: DeadlineExceeded) -> JSONResponse:
    return _error(504, exc.message)


@app.exception_handler(RequestValidationError)
async def _validation_error(request, exc: RequestValidationError) -> JSONResponse:
    # FastAPI's default body echoes the offending input back; keep the status.
    return _error(422, "Invalid request")


class _GuardMiddleware:
    """Refuse oversized bodies up front and turn crashes into a fixed 500.

    Starlette parses the whole multipart body before an endpoint runs, so the
    per-file cap in ``_save_upload`` alone would still spool an oversized
    upload to disk first. Requests without a Content-Length (chunked) are left
    to a fronting proxy's body limit and that per-file cap.

    Unhandled exceptions are logged with their traceback and answered with a
    generic JSON 500 rather than Starlette's plain-text one.
    """

    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http":
            await self.app(scope, receive, send)
            return

        # The request budget starts here, before Starlette reads the upload
        # (FastAPI parses the whole multipart body before the endpoint runs),
        # and counts the time the request already spent in a reverse proxy.
        headers = {name.decode("latin-1").lower(): value.decode("latin-1") for name, value in scope.get("headers", ())}
        scope.setdefault("state", {})["deadline"] = Deadline.from_headers(
            lambda name: headers.get(name.lower()), REQUEST_DEADLINE_SECONDS
        )

        for name, value in scope.get("headers", ()):
            if name == b"content-length":
                try:
                    too_large = int(value) > MAX_FILE_SIZE + _MULTIPART_SLACK
                except ValueError:
                    too_large = False
                if too_large:
                    await _error(413, "File too large")(scope, receive, send)
                    return
                break

        response_started = False

        async def tracking_send(message):
            nonlocal response_started
            if message["type"] == "http.response.start":
                response_started = True
            await send(message)

        try:
            await self.app(scope, receive, tracking_send)
        except DeadlineExceeded:
            # Raised outside an endpoint (e.g. a dependency); same 504.
            if response_started:
                raise
            await _error(504, filo_deadline.MESSAGE)(scope, receive, send)
        except Exception:
            logger.exception("Unhandled error on %s %s", scope.get("method"), scope.get("path"))
            if response_started:
                raise
            await _error(500, "Conversion failed")(scope, receive, send)


app.add_middleware(_GuardMiddleware)

# ----- Office -> PDF (LibreOffice) -----------------------------------------

# Presentation / document formats handled by the ppt-to-pdf endpoint.
PPT_EXTENSIONS = {
    ".ppt",
    ".pptx",
    ".pps",
    ".ppsx",
    ".odp",
    ".doc",
    ".docx",
}

# Spreadsheet formats handled by the xls-to-pdf endpoint. These use the Calc
# PDF export filter so rows/columns render correctly.
XLS_EXTENSIONS = {
    ".xls",
    ".xlsx",
    ".xlsm",
    ".xlsb",
    ".ods",
    ".csv",
}

PDF_EXTENSIONS = {".pdf"}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# LibreOffice has no PDF *import* filter registered for Writer by default: the
# extension is claimed by Draw, so `--convert-to docx` on a PDF exits 0 and
# writes nothing. Naming the Writer import filter explicitly is what makes
# PDF -> DOCX work at all.
PDF_IMPORT_FILTER = "writer_pdf_import"

# Longest a single LibreOffice run may take before its whole process group is
# killed and the client gets a 504. Each run gets the smaller of this and what
# is left of the request budget (REQUEST_DEADLINE_SECONDS): many reverse
# proxies and CDNs cut a request at 100 s, so a longer run would only burn CPU
# for a client that has already been given up on.
CONVERSION_TIMEOUT_SECONDS = _env_int("LO_TIMEOUT_SECONDS", 85)

# Render resolution for PDF -> PPT page images (higher = sharper but larger).
PDF_TO_PPT_DPI = 150
# Default luminance cut-off (0-255) for black_and_white output.
PDF_TO_PPT_THRESHOLD = 150

SOFFICE_CANDIDATES = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "soffice",
    "libreoffice",
]


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


def get_soffice_path() -> str:
    for candidate in SOFFICE_CANDIDATES:
        if Path(candidate).exists() or shutil.which(candidate):
            return candidate

    logger.error("LibreOffice executable not found (tried %s)", SOFFICE_CANDIDATES)
    raise HTTPException(status_code=500, detail="Conversion failed")


async def _save_upload(file: UploadFile, work_dir: Path, allowed_extensions: set[str]) -> Path:
    """Validate the extension and stream the upload to disk with a size cap."""
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in allowed_extensions:
        raise HTTPException(status_code=400, detail="Unsupported file type")

    input_path = work_dir / f"{uuid4().hex}{suffix}"
    size = 0
    with input_path.open("wb") as target:
        while chunk := await file.read(1024 * 1024):
            size += len(chunk)
            if size > MAX_FILE_SIZE:
                raise HTTPException(status_code=413, detail="File too large")
            target.write(chunk)
    return input_path


def _open_pdf(path: Path) -> "fitz.Document":
    """Open a PDF, reporting unreadable input as 4xx rather than a 500."""
    try:
        pdf = fitz.open(str(path))
    except Exception as exc:
        logger.info("Rejected unreadable PDF: %s", exc)
        raise HTTPException(status_code=400, detail="Invalid or corrupt PDF") from exc

    if pdf.needs_pass:
        pdf.close()
        raise HTTPException(status_code=400, detail="PDF is password protected")
    if pdf.page_count > MAX_PDF_PAGES:
        pdf.close()
        raise HTTPException(status_code=422, detail="PDF has too many pages")
    return pdf


def _open_pdfplumber(path: Path):
    """pdfplumber equivalent of :func:`_open_pdf`."""
    try:
        pdf = pdfplumber.open(str(path))
    except Exception as exc:
        logger.info("Rejected unreadable PDF: %s", exc)
        raise HTTPException(
            status_code=400, detail="Invalid, corrupt, or password-protected PDF"
        ) from exc

    if len(pdf.pages) > MAX_PDF_PAGES:
        pdf.close()
        raise HTTPException(status_code=422, detail="PDF has too many pages")
    return pdf


def _check_render_size(rect: "fitz.Rect", zoom: float) -> None:
    """Refuse to rasterise a region whose pixmap would exceed MAX_RENDER_PIXELS."""
    if abs(rect.width * zoom) * abs(rect.height * zoom) > MAX_RENDER_PIXELS:
        raise HTTPException(status_code=422, detail="PDF page is too large")


def _watermark(request: Request) -> bool:
    return filo_watermark.watermark_requested(request.headers.get(filo_watermark.HEADER))


def _file_response(result_path: Path, work_dir: Path, download_name: str, media_type: str) -> FileResponse:
    """Return a FileResponse that cleans up ``work_dir`` after streaming."""
    return FileResponse(
        result_path,
        media_type=media_type,
        filename=download_name,
        background=BackgroundTask(shutil.rmtree, work_dir, ignore_errors=True),
    )


# ----- LibreOffice runner ---------------------------------------------------


# Hardened settings copied into every LibreOffice profile before first start:
# macros off, external links never updated, no recalculation on load.
# The key names and their sources are documented in the file itself.
_LO_SETTINGS = Path(__file__).with_name("lo_profile") / "registrymodifications.xcu"
_OOR = "{http://openoffice.org/2001/registry}"

# A user profile built from scratch costs LibreOffice 1-2s of first-run setup on
# every conversion. Seeding one template and copying it (~0.5 MB, a few ms) buys
# that back while still handing each run a private profile. The template is only
# fully populated once a real conversion has exercised the import/export
# filters; --terminate_after_init leaves it half-built and no faster.
_profile_template_dir: Path | None = None
_profile_template_guard = threading.Lock()


# Environment variable names that look like credentials; stripped before
# soffice starts (see _soffice_env).
_SECRET_ENV = re.compile(r"SECRET|TOKEN|PASSWORD|PASSWD|API_KEY|PRIVATE_KEY|_KEYS?$", re.IGNORECASE)


@lru_cache(maxsize=1)
def _soffice_env() -> dict[str, str]:
    """Environment for soffice, with every writable path under the temp dir.

    The container runs with a read-only root filesystem, so LibreOffice (and
    the fontconfig cache it may touch) must never fall back to a HOME it
    cannot write. Profiles are passed explicitly via -env:UserInstallation.
    """
    # Never hand deployment secrets to a process parsing untrusted files.
    env = {key: value for key, value in os.environ.items() if not _SECRET_ENV.search(key)}
    scratch = tempfile.gettempdir()
    home = env.get("HOME", "")
    if not home or not os.access(home, os.W_OK):
        env["HOME"] = scratch
    env.setdefault("TMPDIR", scratch)
    return env


def _settings(path: Path) -> set[tuple[str, str, str]]:
    """(item path, prop name, value) triples from a registrymodifications.xcu."""
    triples = set()
    for item in ElementTree.parse(path).getroot().iter("item"):
        for prop in item.iter("prop"):
            value = prop.find("value")
            triples.add(
                (
                    item.get(f"{_OOR}path", ""),
                    prop.get(f"{_OOR}name", ""),
                    "" if value is None else (value.text or ""),
                )
            )
    return triples


def _seed_profile(profile: Path) -> None:
    """Install the hardened settings into a not-yet-started profile.

    Raises if the settings file is missing: running LibreOffice without them
    would silently re-enable macros, so failing the conversion is preferable.
    """
    user_dir = profile / "user"
    user_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(_LO_SETTINGS, user_dir / "registrymodifications.xcu")


def _run_soffice_process(
    command: list[str], timeout: float | None = None, tmp_dir: Path | None = None
) -> subprocess.CompletedProcess:
    """Run soffice in its own process group, killing the group on timeout.

    ``soffice`` is a launcher that forks soffice.bin; killing only the direct
    child on timeout (what ``subprocess.run`` does) leaves the real worker
    running and holding memory and a PID slot.

    ``tmp_dir`` (the request's work dir) receives LibreOffice's own scratch
    files, so they go with the request's cleanup. Its IPC pipe
    (OSL_PIPE_*) is always created in /tmp: a killed run leaves that empty
    socket behind in /tmp.
    """
    env = _soffice_env()
    if tmp_dir is not None:
        tmp_dir.mkdir(parents=True, exist_ok=True)
        env = {**env, "TMPDIR": str(tmp_dir)}
    process = subprocess.Popen(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        errors="replace",
        env=env,
        start_new_session=True,
    )
    try:
        stdout, stderr = process.communicate(
            timeout=CONVERSION_TIMEOUT_SECONDS if timeout is None else timeout
        )
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.communicate()
        raise
    return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)


def _build_profile_template() -> Path:
    """Populate a profile by converting a throwaway one-glyph PDF."""
    root = Path(tempfile.mkdtemp(prefix="lo_template_"))
    atexit.register(shutil.rmtree, root, ignore_errors=True)

    profile = root / "profile"
    seed_dir = root / "seed"
    seed_dir.mkdir()
    seed_pdf = seed_dir / "seed.pdf"
    _seed_profile(profile)

    with fitz.open() as document:
        page = document.new_page(width=10, height=10)
        page.insert_text((2, 8), "x", fontsize=6)
        document.save(str(seed_pdf))

    _run_soffice_process(
        [
            get_soffice_path(),
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nodefault",
            "--nolockcheck",
            f"-env:UserInstallation=file://{profile}",
            f"--infilter={PDF_IMPORT_FILTER}",
            "--convert-to",
            "docx:MS Word 2007 XML",
            "--outdir",
            str(seed_dir),
            str(seed_pdf),
        ]
    )

    # LibreOffice rewrites registrymodifications.xcu on exit. It keeps user
    # values, but verify rather than trust it: every copy of this template
    # inherits whatever ends up in that file.
    written = profile / "user" / "registrymodifications.xcu"
    try:
        hardened = _settings(_LO_SETTINGS) <= _settings(written)
    except Exception:
        hardened = False
    if not hardened:
        logger.warning("LibreOffice dropped hardened settings from the template; re-seeding")
        _seed_profile(profile)
    return profile


def _profile_template() -> Path | None:
    """Return the shared profile template, building it once on first use."""
    global _profile_template_dir

    with _profile_template_guard:
        if _profile_template_dir is None or not _profile_template_dir.exists():
            try:
                _profile_template_dir = _build_profile_template()
            except Exception:
                # Fall back to letting LibreOffice build a profile per run: a
                # slower conversion beats a failed one.
                logger.warning("LibreOffice profile template build failed", exc_info=True)
                _profile_template_dir = None
        return _profile_template_dir


def _run_soffice(
    input_path: Path,
    output_dir: Path,
    convert_to: str,
    input_filter: str | None = None,
    deadline: Deadline | None = None,
) -> None:
    """Convert a file with LibreOffice headless, raising on a failed run.

    The run is killed (504) after min(LO_TIMEOUT_SECONDS, what is left of the
    request ``deadline``).

    Every call gets a throw-away user profile. LibreOffice refuses to start a
    second instance against a profile that is already locked and quietly exits
    0 without writing output, so concurrent requests would otherwise return
    "file was not created" at random.

    Blocking: callers run it in a worker thread while holding ``_HEAVY_SLOTS``.
    """
    profile_dir = output_dir.parent / "lo_profile"
    template = _profile_template()
    if template is not None:
        shutil.copytree(template, profile_dir)
    else:
        _seed_profile(profile_dir)

    command = [
        get_soffice_path(),
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--nodefault",
        "--nolockcheck",
        f"-env:UserInstallation=file://{profile_dir}",
    ]
    if input_filter:
        command.append(f"--infilter={input_filter}")
    command += [
        "--convert-to",
        convert_to,
        "--outdir",
        str(output_dir),
        str(input_path),
    ]

    timeout = CONVERSION_TIMEOUT_SECONDS
    if deadline is not None:
        timeout = deadline.timeout(CONVERSION_TIMEOUT_SECONDS)
    try:
        completed = _run_soffice_process(command, timeout, output_dir.parent / "lo_tmp")
    except subprocess.TimeoutExpired:
        logger.warning("LibreOffice run killed after %.1f s", timeout)
        if deadline is not None:
            deadline.expire()
        raise DeadlineExceeded() from None

    if completed.returncode != 0:
        logger.error(
            "soffice exited %s: %s",
            completed.returncode,
            (completed.stderr.strip() or completed.stdout.strip())[-2000:],
        )
        raise HTTPException(status_code=500, detail="Conversion failed")


async def _soffice_convert(
    request: Request,
    input_path: Path,
    output_dir: Path,
    convert_to: str,
    input_filter: str | None = None,
) -> None:
    """Run :func:`_run_soffice` off the event loop, within the slot cap and the budget."""
    await _run_bounded(
        _HEAVY_SLOTS, request, _run_soffice, input_path, output_dir, convert_to, input_filter, _deadline(request)
    )


def _check_archive(path: Path) -> None:
    """Refuse zip-based uploads (OOXML/ODF) that would unpack to a huge size.

    Uses the sizes declared in the archive's directory. Anything that is not
    a readable zip is left for LibreOffice to judge, as before.
    """
    try:
        if not zipfile.is_zipfile(path):
            return
        with zipfile.ZipFile(path) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except Exception:
        return
    if unpacked > MAX_UNPACKED_MB * 1024 * 1024:
        logger.info("Rejected archive unpacking to %d bytes", unpacked)
        raise HTTPException(status_code=422, detail="File is too large when unpacked")


# ----- Office -> PDF --------------------------------------------------------


async def office_to_pdf(
    request: Request,
    file: UploadFile,
    allowed_extensions: set[str],
    convert_to: str = "pdf",
) -> FileResponse:
    """Convert an uploaded office file to PDF via LibreOffice headless.

    ``allowed_extensions`` restricts the input types per endpoint, and
    ``convert_to`` selects the LibreOffice output filter (e.g. ``pdf`` or
    ``pdf:calc_pdf_Export``).
    """
    work_dir = Path(tempfile.mkdtemp(prefix="lo_convert_"))
    output_dir = work_dir / "out"
    output_dir.mkdir()
    should_cleanup = True

    try:
        input_path = await _save_upload(file, work_dir, allowed_extensions)
        await run_in_threadpool(_check_archive, input_path)

        await _soffice_convert(request, input_path, output_dir, convert_to)

        pdf_files = sorted(output_dir.glob("*.pdf"))
        if not pdf_files:
            logger.error("soffice exited 0 but wrote no PDF for %s", input_path.suffix)
            raise HTTPException(status_code=500, detail="Conversion failed")

        response_path = work_dir / "converted.pdf"
        shutil.copyfile(pdf_files[0], response_path)
        if _watermark(request):
            await run_in_threadpool(filo_watermark.watermark_pdf, response_path)
        should_cleanup = False

        return _file_response(
            response_path,
            work_dir,
            download_name=f"{Path(file.filename or 'converted').stem}.pdf",
            media_type="application/pdf",
        )

    except subprocess.TimeoutExpired as exc:
        raise DeadlineExceeded() from exc

    finally:
        await file.close()
        if should_cleanup:
            shutil.rmtree(work_dir, ignore_errors=True)


@app.post("/convert/ppt-to-pdf")
async def convert_ppt_to_pdf(request: Request, file: UploadFile = File(...)) -> FileResponse:
    return await office_to_pdf(request, file, PPT_EXTENSIONS)


@app.post("/convert/xls-to-pdf")
async def convert_xls_to_pdf(request: Request, file: UploadFile = File(...)) -> FileResponse:
    return await office_to_pdf(request, file, XLS_EXTENSIONS, convert_to="pdf:calc_pdf_Export")


# ----- PDF -> DOCX (visual fidelity) ----------------------------------------


# LibreOffice writes the imported page into a Writer document that keeps a 1 cm
# print margin, and positions every frame relative to that margin instead of the
# paper edge. The whole page therefore lands 1 cm down and right of where the
# PDF drew it. Zeroing the margins puts the coordinate origin back on the paper
# corner, which is what the frame offsets were measured from.
_PAGE_MARGIN = re.compile(r"<w:pgMar\b[^>]*/>")
_ZERO_MARGIN = (
    '<w:pgMar w:left="0" w:right="0" w:gutter="0"'
    ' w:header="0" w:top="0" w:footer="0" w:bottom="0"/>'
)

# One shape as LibreOffice writes it: either a modern DrawingML anchor (wrapped
# in AlternateContent alongside a legacy fallback) or a bare VML <w:pict>.
# Scanning left to right, a pict nested inside an AlternateContent fallback is
# consumed by the first alternative, so only top-level shapes are matched.
_SHAPE = re.compile(
    r"<mc:AlternateContent\b.*?</mc:AlternateContent>|<w:pict\b.*?</w:pict>", re.S
)
_VML_STYLE = re.compile(r'(<v:(?:shape|rect|oval|line|polyline)\b[^>]*?\bstyle=")([^"]*)(")')
_Z_INDEX = re.compile(r";?\s*z-index:\s*-?\d+")
_RELATIVE_HEIGHT = re.compile(r'relativeHeight="\d+"')
_POSITION_H = re.compile(r"<wp:positionH.*?<wp:posOffset>(-?\d+)</wp:posOffset>", re.S)
_POSITION_V = re.compile(r"<wp:positionV.*?<wp:posOffset>(-?\d+)</wp:posOffset>", re.S)
_EXTENT = re.compile(r'<wp:extent cx="(\d+)" cy="(\d+)"/>')
_VML_LENGTH = re.compile(r"([\w-]+):\s*(-?[\d.]+)pt")

# Word stacks drawings in two layers: behind the text and in front of it. The
# bands keep demoted fills in the first and text frames in the second, while the
# running counter preserves paint order within each band.
_BACKGROUND_BAND = -100_000
_FOREGROUND_BAND = 100_000

_EMU_PER_POINT = 12700

# Share of a text frame that must sit under a fill before the fill is treated as
# that frame's backdrop rather than a neighbouring graphic.
_COVERED_RATIO = 0.5


def _drawingml_box(shape: str) -> tuple[float, float, float, float] | None:
    """Bounding box of a DrawingML anchor, in EMU."""
    left = _POSITION_H.search(shape)
    top = _POSITION_V.search(shape)
    extent = _EXTENT.search(shape)
    if not (left and top and extent):
        return None
    x, y = int(left.group(1)), int(top.group(1))
    return (x, y, x + int(extent.group(1)), y + int(extent.group(2)))


def _vml_box(shape: str) -> tuple[float, float, float, float] | None:
    """Bounding box of a VML shape, converted from its CSS style to EMU."""
    style = _VML_STYLE.search(shape)
    if not style:
        return None
    lengths = dict(_VML_LENGTH.findall(style.group(2)))
    try:
        x = float(lengths["margin-left"]) * _EMU_PER_POINT
        y = float(lengths["margin-top"]) * _EMU_PER_POINT
        width = float(lengths["width"]) * _EMU_PER_POINT
        height = float(lengths["height"]) * _EMU_PER_POINT
    except KeyError:
        return None
    return (x, y, x + width, y + height)


def _covers(fill: tuple | None, frame: tuple | None) -> bool:
    """True when ``fill`` hides most of ``frame``."""
    if not (fill and frame):
        return False
    overlap_x = min(fill[2], frame[2]) - max(fill[0], frame[0])
    overlap_y = min(fill[3], frame[3]) - max(fill[1], frame[1])
    if overlap_x <= 0 or overlap_y <= 0:
        return False
    area = (frame[2] - frame[0]) * (frame[3] - frame[1])
    return area > 0 and (overlap_x * overlap_y) / area > _COVERED_RATIO


def _restack_shapes(xml: str) -> str:
    """Drop fills that hide a later text frame into the behind-text layer.

    LibreOffice exports vector fills as legacy VML and text as DrawingML
    frames, and renderers put the legacy layer on top whatever the document
    order or z-index says — so a heading drawn over a coloured banner in the
    PDF ends up hidden underneath it. Giving the offending fill a negative
    z-index is what moves it back down; the check stays geometric because a
    fill that covers text is not always a mistake.

    Text frames also get ``wrap="none"``: each frame holds a single run
    measured from the PDF, so a substituted font that renders a hair wider
    would otherwise wrap and push half the line out of the frame's fixed box.
    """
    shapes = [match.group(0) for match in _SHAPE.finditer(xml)]
    frames = [
        (index, _drawingml_box(shape))
        for index, shape in enumerate(shapes)
        if "txbxContent" in shape
    ]

    # A fill only needs demoting when it hides a frame the PDF drew *later*.
    # Anything else keeps the stacking LibreOffice gave it — a scan laid over
    # the page's own text layer, for instance, is meant to cover it.
    demoted = {
        index
        for index, shape in enumerate(shapes)
        if shape.startswith("<w:pict")
        and any(
            other > index and _covers(_vml_box(shape), box) for other, box in frames
        )
    }

    counter = -1

    def restack(match: re.Match) -> str:
        nonlocal counter
        counter += 1
        shape = match.group(0)

        if "txbxContent" in shape:
            shape = shape.replace('behindDoc="1"', 'behindDoc="0"')
            shape = shape.replace('<wps:bodyPr wrap="square"', '<wps:bodyPr wrap="none"')
            depth = _FOREGROUND_BAND + counter
            return _RELATIVE_HEIGHT.sub(f'relativeHeight="{depth}"', shape)

        if counter not in demoted:
            return shape

        depth = _BACKGROUND_BAND + counter

        def restyle(style: re.Match) -> str:
            declarations = _Z_INDEX.sub("", style.group(2)).strip().rstrip(";")
            return f"{style.group(1)}{declarations};z-index:{depth}{style.group(3)}"

        return _VML_STYLE.sub(restyle, shape)

    return _SHAPE.sub(restack, xml)


def _restore_page_fidelity(docx_path: Path) -> None:
    """Rewrite the .docx in place so it renders like the source PDF.

    A .docx is a zip; only ``word/document.xml`` changes, so every other part
    is copied across verbatim.
    """
    patched_path = docx_path.with_suffix(".patched.docx")

    with zipfile.ZipFile(docx_path) as source:
        with zipfile.ZipFile(patched_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/document.xml":
                    xml = _PAGE_MARGIN.sub(_ZERO_MARGIN, data.decode("utf-8"))
                    data = _restack_shapes(xml).encode("utf-8")
                target.writestr(item, data)

    patched_path.replace(docx_path)


@app.post("/convert/pdf-to-doc")
async def convert_pdf_to_doc(request: Request, file: UploadFile = File(...)) -> FileResponse:
    """Convert a PDF to .docx with LibreOffice, preserving the page's look.

    LibreOffice reconstructs each page as positioned text frames and drawing
    shapes, so the result opens looking like the original — same fonts, same
    coordinates, same graphics. That fidelity is the trade-off: the text does
    not reflow like a hand-written document, since the PDF carries no paragraph
    or table structure to recover.
    """
    work_dir = Path(tempfile.mkdtemp(prefix="pdf2doc_"))
    output_dir = work_dir / "out"
    output_dir.mkdir()
    should_cleanup = True

    try:
        input_path = await _save_upload(file, work_dir, PDF_EXTENSIONS)

        # Reject unreadable/locked/overlong PDFs here: LibreOffice would
        # otherwise stall on a password prompt or report an opaque filter error.
        await run_in_threadpool(lambda: _open_pdf(input_path).close())

        await _soffice_convert(
            request,
            input_path,
            output_dir,
            "docx:MS Word 2007 XML",
            PDF_IMPORT_FILTER,
        )

        docx_files = sorted(output_dir.glob("*.docx"))
        if not docx_files:
            logger.error("soffice exited 0 but wrote no DOCX")
            raise HTTPException(status_code=500, detail="Conversion failed")

        response_path = work_dir / "converted.docx"
        shutil.copyfile(docx_files[0], response_path)
        await run_in_threadpool(_restore_page_fidelity, response_path)
        if _watermark(request):
            await run_in_threadpool(filo_watermark.watermark_docx, response_path)
        should_cleanup = False

        return _file_response(
            response_path,
            work_dir,
            download_name=f"{Path(file.filename or 'converted').stem}.docx",
            media_type=DOCX_MIME,
        )

    except subprocess.TimeoutExpired as exc:
        raise DeadlineExceeded() from exc

    finally:
        await file.close()
        if should_cleanup:
            shutil.rmtree(work_dir, ignore_errors=True)


# ----- PDF -> Excel (table extraction) --------------------------------------


# Table-detection passes, tried in order. Ruled tables (visible borders) are
# most accurate; the text strategy recovers borderless / whitespace-aligned
# tables that the line strategy misses.
_TABLE_STRATEGIES = (
    {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
    {"vertical_strategy": "text", "horizontal_strategy": "text"},
)


# A real table is mostly populated. Chart gridlines and decorative boxes get
# detected as sparse "tables" holding a few stray axis/data labels, so anything
# below this fill ratio is rejected in favour of plain text extraction.
_MIN_TABLE_FILL_RATIO = 0.6


# Share of adjacent cell pairs that may split a word before the "table" is
# treated as prose sliced up by chart gridlines rather than real columns.
_MAX_WORD_SPLIT_RATIO = 0.3


def _is_chopped_prose(rows: list[list]) -> bool:
    """True when cell boundaries cut through words instead of between values.

    Ruling lines from charts slice body text into fake columns, producing cells
    like "Maecenas eu l" | "igula ege". Genuine table cells hold whole values,
    so a lowercase letter running straight into the next cell is a strong hint
    the detection is bogus.
    """
    pairs = splits = 0
    for row in rows:
        for left, right in zip(row, row[1:]):
            if not (left and right):
                continue
            pairs += 1
            if left[-1].islower() and right[0].islower():
                splits += 1
    return bool(pairs) and splits / pairs > _MAX_WORD_SPLIT_RATIO


def _looks_like_table(rows: list[list]) -> bool:
    """Reject sparse detections (chart gridlines, decorative rectangles)."""
    if len(rows) < 2 or len(rows[0]) < 2:
        # A single row or column is a misread that drops data; plain text wins.
        return False
    total = sum(len(row) for row in rows)
    filled = sum(1 for row in rows for cell in row if cell)
    if not total or filled / total < _MIN_TABLE_FILL_RATIO:
        return False
    # Every row of a genuine table should carry at least two populated cells.
    if not all(sum(1 for cell in row if cell) >= 2 for row in rows):
        return False
    return not _is_chopped_prose(rows)


def _extract_page_tables(page) -> list[list[list]]:
    """Return the best set of tables for a page, trying strategies in order."""
    for settings in _TABLE_STRATEGIES:
        try:
            tables = page.extract_tables(table_settings=settings)
        except Exception:
            tables = []
        cleaned = [
            t for t in (_clean_table(t) for t in tables) if _looks_like_table(t)
        ]
        if cleaned:
            return cleaned
    return []


def _clean_table(table) -> list[list]:
    """Normalize cells and drop fully-empty rows/columns."""
    if not table:
        return []
    rows = [["" if c is None else str(c).strip() for c in row] for row in table]
    rows = [r for r in rows if any(cell for cell in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    return [[r[i] for i in keep] for r in rows]


# Visuals smaller than this (in points) are page furniture — rules, underlines,
# cell borders — rather than a real picture or chart.
_MIN_VISUAL_POINTS = 40
# Rendering zoom for extracted visuals; 2x (144 DPI) keeps chart labels legible.
_VISUAL_ZOOM = 2.0
# Excel's column width unit is roughly one character ~= 7 px; row height is in
# points, and 1 px = 0.75 pt.
_PX_PER_WIDTH_UNIT = 7.0
_PT_PER_PX = 0.75


# A fill spanning most of the page width but only a line or two tall is a text
# background, not artwork. Leaving these in makes every cluster swallow the
# surrounding prose.
_TEXT_BAND_WIDTH_RATIO = 0.6


# Overlap (as a share of the smaller region) above which two visuals are the
# same artwork picked up twice.
_DUPLICATE_OVERLAP = 0.8


def _overlap_ratio(a: "fitz.Rect", b: "fitz.Rect") -> float:
    """Intersection area as a fraction of the smaller rectangle."""
    intersection = a & b
    if intersection.is_empty:
        return 0.0
    smaller = min(abs(a.get_area()), abs(b.get_area()))
    return abs(intersection.get_area()) / smaller if smaller else 0.0


def _chart_regions(page) -> list["fitz.Rect"]:
    """Cluster vector paths into chart-sized rectangles."""
    page_width = page.rect.width
    artwork = [
        drawing
        for drawing in page.get_drawings()
        if not (
            drawing["rect"].width > _TEXT_BAND_WIDTH_RATIO * page_width
            and drawing["rect"].height < _MIN_VISUAL_POINTS
        )
    ]
    if not artwork:
        return []
    return [
        rect
        for rect in page.cluster_drawings(drawings=artwork)
        if rect.width >= _MIN_VISUAL_POINTS and rect.height >= _MIN_VISUAL_POINTS
    ]


def _visual_regions(page) -> list[tuple[str, "fitz.Rect"]]:
    """Locate raster pictures and vector chart clusters on a page.

    Charts are drawn as vector paths, so ``get_images`` cannot see them —
    ``cluster_drawings`` groups those paths into chart-sized rectangles. Both
    kinds are returned as clip rectangles in reading order.
    """
    regions: list[tuple[str, fitz.Rect]] = []

    for xref, *_ in page.get_images(full=True):
        regions.extend(("image", rect) for rect in page.get_image_rects(xref))

    images = [rect for kind, rect in regions if kind == "image"]

    # Pictures often sit on a background/frame path, which clusters into a
    # near-identical rectangle. Keep the picture and drop the duplicate.
    regions.extend(
        ("chart", rect)
        for rect in _chart_regions(page)
        if not any(_overlap_ratio(rect, img) > _DUPLICATE_OVERLAP for img in images)
    )

    regions.sort(key=lambda item: (round(item[1].y0), item[1].x0))
    return regions


def _extract_visuals(
    pdf_path: Path,
    work_dir: Path,
    workbook,
    max_width_px: int = 700,
    deadline: Deadline | None = None,
) -> int:
    """Embed pictures and charts into a dedicated 'Visuals' sheet.

    Each visual is rendered from its page region, so vector charts come through
    as well as embedded photos. Rows and the picture column are sized to the
    rendered dimensions so nothing overlaps.
    """
    sheet = None
    count = 0
    row = 1
    widest_px = 0
    matrix = fitz.Matrix(_VISUAL_ZOOM, _VISUAL_ZOOM)

    with fitz.open(str(pdf_path)) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            if deadline is not None:
                deadline.check()
            for kind, rect in _visual_regions(page):
                _check_render_size(rect, _VISUAL_ZOOM)
                pixmap = page.get_pixmap(matrix=matrix, clip=rect)
                if not pixmap.width or not pixmap.height:
                    continue

                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                if image.width > max_width_px:
                    ratio = max_width_px / image.width
                    image = image.resize(
                        (max_width_px, max(1, int(image.height * ratio))), Image.LANCZOS
                    )

                image_path = work_dir / f"{kind}_p{page_number}_{count}.png"
                image.save(image_path)

                if sheet is None:
                    sheet = workbook.create_sheet(title="Visuals")
                sheet.cell(row=row, column=1, value=f"Page {page_number}")
                sheet.cell(row=row, column=2, value=kind)
                sheet.add_image(XLImage(str(image_path)), f"C{row}")

                # Size the row to the picture so consecutive visuals never overlap.
                sheet.row_dimensions[row].height = image.height * _PT_PER_PX
                widest_px = max(widest_px, image.width)
                row += 1
                count += 1

    if sheet is not None:
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 8
        sheet.column_dimensions["C"].width = widest_px / _PX_PER_WIDTH_UNIT

    return count


def _build_workbook(
    input_path: Path,
    work_dir: Path,
    include_images: bool,
    watermark: bool = False,
    deadline: Deadline | None = None,
) -> Path:
    """Blocking body of :func:`convert_pdf_to_xls`; runs in a worker thread.

    Stops (DeadlineExceeded) at the next page once ``deadline`` is spent.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    produced_any = False
    found_table = False
    header: list[str] | None = None

    with _open_pdfplumber(input_path) as pdf:
        for page in pdf.pages:
            if deadline is not None:
                deadline.check()
            tables = _extract_page_tables(page)
            if tables:
                found_table = True
                for table in tables:
                    rows = table
                    # Remember the first table's header; drop it when the
                    # same row reappears at the top of a later page.
                    if header is None:
                        header = rows[0]
                    elif rows[0] == header:
                        rows = rows[1:]
                    for row in rows:
                        sheet.append(row)
                    produced_any = True
                continue

            # Fallback: keep the page's text so nothing is silently dropped.
            text = page.extract_text() or ""
            for line in text.splitlines():
                if line.strip():
                    sheet.append([line])
                    produced_any = True

    _autosize_columns(sheet)

    if include_images and _extract_visuals(input_path, work_dir, workbook, deadline=deadline):
        produced_any = True
        # With no real table, 'Data' is just loose prose — land the reader
        # on the artwork instead, which is the useful content.
        if not found_table:
            workbook.active = workbook["Visuals"]

    if not produced_any:
        raise HTTPException(
            status_code=422,
            detail="No extractable text found — the PDF appears to be scanned or image-only (OCR required)",
        )

    if watermark:
        filo_watermark.watermark_workbook(workbook)
    response_path = work_dir / "converted.xlsx"
    workbook.save(str(response_path))
    return response_path


@app.post("/convert/pdf-to-xls")
async def convert_pdf_to_xls(
    request: Request,
    file: UploadFile = File(...),
    include_images: bool = Form(True),
) -> FileResponse:
    """Extract tabular data from a (digital, non-scanned) PDF into an .xlsx.

    Every page is appended to a single continuous 'Data' sheet, so a table
    split across pages comes out as one dataset. A header row repeated on
    continuation pages is written only once.

    Detection is attempted with ruled-line and then text-based strategies;
    sparse detections (chart gridlines, decorative boxes) are rejected. If a
    page yields no table, its text is still written (one line per row) so
    nothing is silently dropped. With ``include_images`` any pictures and
    vector charts are rendered onto a separate 'Visuals' sheet, sized to their
    own dimensions. Scanned PDFs with no text and no visuals return 422.
    """
    work_dir = Path(tempfile.mkdtemp(prefix="pdf2xls_"))
    should_cleanup = True

    try:
        input_path = await _save_upload(file, work_dir, PDF_EXTENSIONS)

        response_path = await _run_bounded(
            _PDF_SLOTS,
            request,
            _build_workbook,
            input_path,
            work_dir,
            include_images,
            _watermark(request),
            _deadline(request),
        )
        should_cleanup = False

        return _file_response(
            response_path,
            work_dir,
            download_name=f"{Path(file.filename or 'converted').stem}.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    finally:
        await file.close()
        if should_cleanup:
            shutil.rmtree(work_dir, ignore_errors=True)


def _autosize_columns(sheet, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(width + 2, max_width)


# ----- PDF -> PPT (one image slide per page) --------------------------------


def _process_pixmap(
    pixmap: "fitz.Pixmap",
    grayscale: bool,
    black_and_white: bool,
    threshold: int,
    invert: bool,
) -> io.BytesIO:
    """Turn a rendered page into an in-memory PNG, applying optional filters.

    When no filters are requested the pixmap is encoded straight to PNG (no
    Pillow round-trip). Threshold/grayscale use Pillow's C-level ``point`` /
    ``convert`` rather than per-pixel Python loops, so they stay fast.
    """
    buffer = io.BytesIO()

    if not (grayscale or black_and_white or invert):
        buffer.write(pixmap.tobytes("png"))
        buffer.seek(0)
        return buffer

    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)

    if grayscale or black_and_white:
        image = image.convert("L")
        if black_and_white:
            # 1-bit output keeps black/white slides tiny.
            image = image.point(lambda p: 255 if p > threshold else 0).convert("1")

    if invert:
        # ImageOps.invert cannot handle 1-bit images; flip via point instead.
        image = image.point(lambda p: 255 - p) if image.mode == "1" else PIL.ImageOps.invert(image)

    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _build_presentation(
    input_path: Path,
    work_dir: Path,
    grayscale: bool,
    black_and_white: bool,
    threshold: int,
    invert: bool,
    watermark: bool = False,
    deadline: Deadline | None = None,
) -> Path:
    """Blocking body of :func:`convert_pdf_to_ppt`; runs in a worker thread.

    Stops (DeadlineExceeded) at the next page once ``deadline`` is spent.
    """
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    zoom = PDF_TO_PPT_DPI / 72.0
    matrix = fitz.Matrix(zoom, zoom)

    with _open_pdf(input_path) as pdf:
        if pdf.page_count == 0:
            raise HTTPException(status_code=422, detail="PDF has no pages")

        for page_index, page in enumerate(pdf):
            if deadline is not None:
                deadline.check()
            _check_render_size(page.rect, zoom)
            pixmap = page.get_pixmap(matrix=matrix)
            image_stream = _process_pixmap(
                pixmap, grayscale, black_and_white, threshold, invert
            )

            # python-pptx uses one slide size for the whole deck; derive it
            # from the first page (typical PDFs have a uniform page size).
            if page_index == 0:
                presentation.slide_width = Emu(int(pixmap.width / PDF_TO_PPT_DPI * 914400))
                presentation.slide_height = Emu(int(pixmap.height / PDF_TO_PPT_DPI * 914400))

            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                image_stream,
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

    if watermark:
        filo_watermark.watermark_presentation(presentation)
    response_path = work_dir / "converted.pptx"
    presentation.save(str(response_path))
    return response_path


@app.post("/convert/pdf-to-ppt")
async def convert_pdf_to_ppt(
    request: Request,
    file: UploadFile = File(...),
    grayscale: bool = Form(False),
    black_and_white: bool = Form(False),
    threshold: int = Form(PDF_TO_PPT_THRESHOLD),
    invert: bool = Form(False),
) -> FileResponse:
    """Render each PDF page to an image and place it on its own slide.

    Preserves the exact visual layout of the PDF (image slides, not editable
    text). Optional image filters:

    * ``grayscale`` — convert pages to greyscale.
    * ``black_and_white`` — 1-bit output using ``threshold`` (0-255); great for
      shrinking scanned text decks.
    * ``invert`` — invert colors (e.g. for dark-mode slides).
    """
    if not 0 <= threshold <= 255:
        raise HTTPException(status_code=400, detail="threshold must be between 0 and 255")

    work_dir = Path(tempfile.mkdtemp(prefix="pdf2ppt_"))
    should_cleanup = True

    try:
        input_path = await _save_upload(file, work_dir, PDF_EXTENSIONS)

        response_path = await _run_bounded(
            _PDF_SLOTS,
            request,
            _build_presentation,
            input_path,
            work_dir,
            grayscale,
            black_and_white,
            threshold,
            invert,
            _watermark(request),
            _deadline(request),
        )
        should_cleanup = False

        return _file_response(
            response_path,
            work_dir,
            download_name=f"{Path(file.filename or 'converted').stem}.pptx",
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

    finally:
        await file.close()
        if should_cleanup:
            shutil.rmtree(work_dir, ignore_errors=True)


# ----- PDF -> text ----------------------------------------------------------

# Page separator written when include_page_markers is on. Kept distinctive so
# downstream consumers can split the file back into pages.
_PAGE_MARKER = "----- Page {number} -----"


def _page_text(page, flags: int, sort_reading_order: bool) -> str:
    """Extract one page's text, trimming trailing whitespace per line."""
    text = page.get_text("text", flags=flags, sort=sort_reading_order)
    return _tidy(text)


def _tidy(text: str) -> str:
    """Strip trailing spaces and collapse runs of blank lines."""
    text = "\n".join(line.rstrip() for line in text.splitlines())
    # Sorted/OCR extraction pads blocks apart; keep a paragraph break to one
    # blank line.
    return re.sub(r"\n{3,}", "\n\n", text)


# OCR renders at a higher resolution than display: Tesseract's accuracy falls
# off sharply below ~300 DPI.
OCR_DPI = 300
OCR_MODES = ("auto", "force", "off")

# Tesseract runs as a subprocess, so threads give real parallelism across pages.
# Four is where the gain flattens on an 8-core host, and it leaves cores free
# for other requests instead of letting one document monopolise the machine.
OCR_WORKERS = max(1, min(4, os.cpu_count() or 1))


def _require_tesseract() -> None:
    if shutil.which("tesseract") is None:
        logger.error("OCR requested but the tesseract binary is not installed")
        raise HTTPException(status_code=503, detail="OCR unavailable")


@lru_cache(maxsize=1)
def _ocr_languages() -> tuple[str, ...]:
    """Installed Tesseract languages.

    Cached because the probe shells out to tesseract, and language packs do not
    come and go while the service is running.
    """
    try:
        return tuple(pytesseract.get_languages(config=""))
    except Exception:
        return ()


def _has_ink(page) -> bool:
    """True if a page holds artwork worth rasterising for OCR.

    A textless page with no images or vector paths is simply blank, and OCR of
    blank paper costs a rasterise + Tesseract pass to return nothing.
    """
    return bool(page.get_images() or page.get_drawings())


GARBLED_MIN_CHARS = 30
GARBLED_RATIO = 0.35


def _is_garbled_char(char: str) -> bool:
    code = ord(char)
    return 0xA0 <= code <= 0xFF or 0xE000 <= code <= 0xF8FF or code == 0xFFFD


def _has_garbled_text(page) -> bool:
    totals: dict[str, int] = {}
    garbled: dict[str, int] = {}
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                font = span["font"]
                for char in span["text"]:
                    if char.isspace() or char.isdigit():
                        continue
                    totals[font] = totals.get(font, 0) + 1
                    if _is_garbled_char(char):
                        garbled[font] = garbled.get(font, 0) + 1
    return any(
        garbled.get(font, 0) >= GARBLED_MIN_CHARS
        and garbled[font] / total >= GARBLED_RATIO
        for font, total in totals.items()
    )


def _render_for_ocr(page) -> Image.Image:
    """Rasterise a page straight into a Pillow image.

    The samples are handed over as-is rather than round-tripped through PNG:
    encoding and decoding cost more than a fifth of a second per page and
    return the very same pixels. Colour is kept — greyscaling is faster still
    but measurably degrades what Tesseract reads back.
    """
    _check_render_size(page.rect, OCR_DPI / 72.0)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(OCR_DPI / 72.0, OCR_DPI / 72.0))
    return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)


def _ocr_image(image: Image.Image, language: str, deadline: Deadline | None = None) -> str:
    """Read one rendered page with Tesseract.

    With a ``deadline``, tesseract is killed when the request budget runs out
    (pytesseract's own timeout) and DeadlineExceeded is raised.
    """
    timeout = 0 if deadline is None else deadline.timeout()
    try:
        return pytesseract.image_to_string(image, lang=language, timeout=timeout)
    # TesseractError is a RuntimeError subclass, so it must be handled first.
    except pytesseract.TesseractError as exc:
        logger.error("Tesseract failed: %s", exc)
        raise HTTPException(status_code=500, detail="OCR failed") from exc
    except RuntimeError as exc:
        # pytesseract's timeout: a plain RuntimeError("Tesseract process timeout").
        if deadline is not None and "timeout" in str(exc).lower():
            deadline.expire()
            raise DeadlineExceeded() from exc
        raise


def _extract_text_layer(
    pdf_path: Path,
    flags: int,
    sort_reading_order: bool,
    ocr: str,
    deadline: Deadline | None = None,
) -> tuple[list[str], list[int]]:
    """Return each page's text layer, plus the pages that still need OCR.

    Splitting this from :func:`_ocr_pages` lets digital PDFs (the common case)
    skip the OCR concurrency slot entirely; only documents that really need
    Tesseract queue for it.

    Runs synchronously; the caller hands it to a worker thread rather than
    stalling the event loop.
    """
    pages: list[str] = []
    needs_ocr: list[int] = []

    with _open_pdf(pdf_path) as pdf:
        for index, page in enumerate(pdf):
            if deadline is not None:
                deadline.check()
            if ocr == "force":
                pages.append("")
                needs_ocr.append(index)
                continue

            text = _page_text(page, flags, sort_reading_order)
            # A page with no text layer is either blank or a scan; only the
            # latter is worth the cost of rasterising and reading.
            if ocr == "auto" and (
                (not text.strip() and _has_ink(page)) or _has_garbled_text(page)
            ):
                pages.append("")
                needs_ocr.append(index)
            else:
                pages.append(text)

    return pages, needs_ocr


def _ocr_pages(
    pdf_path: Path,
    pages: list[str],
    needs_ocr: list[int],
    ocr_language: str,
    deadline: Deadline | None = None,
) -> None:
    """OCR the ``needs_ocr`` pages, writing their text into ``pages`` in place.

    Pages are read concurrently, which is where nearly all the time goes on a
    scanned document. They are rasterised a batch at a time so at most
    ``OCR_WORKERS`` full-page images — around 26 MB each at 300 DPI — are ever
    held in memory at once.

    Runs synchronously; the caller hands it to a worker thread while holding
    ``_HEAVY_SLOTS``.
    """
    with _open_pdf(pdf_path) as pdf:
        with ThreadPoolExecutor(max_workers=OCR_WORKERS) as pool:
            for start in range(0, len(needs_ocr), OCR_WORKERS):
                batch = needs_ocr[start : start + OCR_WORKERS]
                images = []
                for index in batch:
                    if deadline is not None:
                        deadline.check()
                    images.append(_render_for_ocr(pdf[index]))
                texts = pool.map(lambda image: _ocr_image(image, ocr_language, deadline), images)
                for index, text in zip(batch, texts):
                    pages[index] = _tidy(text)


@app.post("/convert/pdf-to-text")
async def convert_pdf_to_text(
    request: Request,
    file: UploadFile = File(...),
    include_page_markers: bool = Form(True),
    sort_reading_order: bool = Form(False),
    dehyphenate: bool = Form(False),
    ocr: str = Form("auto"),
    ocr_language: str = Form("eng"),
) -> FileResponse:
    """Extract a PDF's text to a UTF-8 .txt file, with OCR for scanned pages.

    Options:

    * ``include_page_markers`` — insert a ``----- Page N -----`` separator so
      the output can be split back into pages.
    * ``sort_reading_order`` — re-order blocks top-left to bottom-right. Fixes
      PDFs whose internal text order is scrambled, but also indents text to
      mirror the page layout, so it is off by default.
    * ``dehyphenate`` — rejoin words split across line breaks.
    * ``ocr`` — ``auto`` reads any page lacking a text layer with Tesseract,
      leaving digital pages untouched; ``force`` OCRs every page (use when the
      embedded text layer is itself garbage); ``off`` never OCRs.
    * ``ocr_language`` — Tesseract language code, e.g. ``eng`` or ``eng+deu``.

    OCR needs the tesseract binary; without it an OCR request returns 503.
    """
    if ocr not in OCR_MODES:
        raise HTTPException(status_code=400, detail="ocr must be one of auto, force, off")

    if ocr != "off":
        _require_tesseract()
        available = await run_in_threadpool(_ocr_languages)
        requested = ocr_language.split("+")
        if available and not all(code in available for code in requested):
            raise HTTPException(status_code=400, detail="OCR language is not installed")

    work_dir = Path(tempfile.mkdtemp(prefix="pdf2txt_"))
    should_cleanup = True

    flags = fitz.TEXTFLAGS_TEXT
    if dehyphenate:
        flags |= fitz.TEXT_DEHYPHENATE

    try:
        input_path = await _save_upload(file, work_dir, PDF_EXTENSIONS)

        pages, needs_ocr = await _run_bounded(
            _PDF_SLOTS,
            request,
            _extract_text_layer,
            input_path,
            flags,
            sort_reading_order,
            ocr,
            _deadline(request),
        )
        if needs_ocr:
            await _run_bounded(
                _HEAVY_SLOTS,
                request,
                _ocr_pages,
                input_path,
                pages,
                needs_ocr,
                ocr_language,
                _deadline(request),
            )

        if not any(page.strip() for page in pages):
            detail = (
                "No text found even after OCR — the PDF may be blank or unreadable"
                if ocr != "off"
                else "No text layer found — the PDF appears to be scanned or image-only (retry with ocr=auto)"
            )
            raise HTTPException(status_code=422, detail=detail)

        chunks: list[str] = []
        for page_number, text in enumerate(pages, start=1):
            if include_page_markers:
                chunks.append(_PAGE_MARKER.format(number=page_number))
            chunks.append(text.strip("\n"))

        response_path = work_dir / "converted.txt"
        response_path.write_text("\n\n".join(chunks).strip() + "\n", encoding="utf-8")
        should_cleanup = False

        return _file_response(
            response_path,
            work_dir,
            download_name=f"{Path(file.filename or 'converted').stem}.txt",
            media_type="text/plain; charset=utf-8",
        )

    finally:
        await file.close()
        if should_cleanup:
            shutil.rmtree(work_dir, ignore_errors=True)
